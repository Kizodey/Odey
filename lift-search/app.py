import os
import zipfile
import shutil
import uuid
import json
from pathlib import Path
from flask import Flask, request, jsonify, render_template, session
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
import anthropic
import pdfplumber
import docx
import openpyxl

load_dotenv()

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", os.urandom(24))

UPLOAD_FOLDER = Path("uploads")
EXTRACTED_FOLDER = Path("extracted")
MAX_CONTENT_LENGTH = 100 * 1024 * 1024  # 100MB
ALLOWED_EXTENSIONS = {"zip"}

app.config["MAX_CONTENT_LENGTH"] = MAX_CONTENT_LENGTH

client = anthropic.Anthropic(api_key=os.environ.get("ANTHROPIC_API_KEY"))

DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt", ".csv", ".rtf"}


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_text_from_pdf(path: Path) -> str:
    try:
        with pdfplumber.open(path) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as e:
        return f"[Error reading PDF: {e}]"


def extract_text_from_docx(path: Path) -> str:
    try:
        doc = docx.Document(path)
        return "\n".join(para.text for para in doc.paragraphs)
    except Exception as e:
        return f"[Error reading DOCX: {e}]"


def extract_text_from_xlsx(path: Path) -> str:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        parts = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            parts.append(f"[Sheet: {sheet}]")
            for row in ws.iter_rows(values_only=True):
                row_text = "\t".join(str(c) if c is not None else "" for c in row)
                if row_text.strip():
                    parts.append(row_text)
        return "\n".join(parts)
    except Exception as e:
        return f"[Error reading XLSX: {e}]"


def extract_text_from_txt(path: Path) -> str:
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return "[Error: could not decode file]"


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    elif suffix in (".docx", ".doc"):
        return extract_text_from_docx(path)
    elif suffix in (".xlsx", ".xls"):
        return extract_text_from_xlsx(path)
    elif suffix in (".txt", ".csv", ".rtf"):
        return extract_text_from_txt(path)
    return ""


def collect_documents(session_dir: Path) -> list[dict]:
    docs = []
    for path in sorted(session_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in DOCUMENT_EXTENSIONS:
            relative = path.relative_to(session_dir)
            text = extract_text(path)
            if text.strip():
                docs.append({"filename": str(relative), "content": text[:50000]})
    return docs


def build_document_context(docs: list[dict]) -> str:
    parts = []
    for i, doc in enumerate(docs, 1):
        parts.append(f"=== Document {i}: {doc['filename']} ===\n{doc['content']}\n")
    return "\n".join(parts)


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/upload", methods=["POST"])
def upload():
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400

    file = request.files["file"]
    if file.filename == "":
        return jsonify({"error": "No file selected"}), 400

    if not allowed_file(file.filename):
        return jsonify({"error": "Only ZIP files are accepted"}), 400

    session_id = str(uuid.uuid4())
    upload_path = UPLOAD_FOLDER / f"{session_id}.zip"
    extract_dir = EXTRACTED_FOLDER / session_id

    upload_path.parent.mkdir(exist_ok=True)
    extract_dir.mkdir(parents=True, exist_ok=True)

    file.save(upload_path)

    try:
        with zipfile.ZipFile(upload_path, "r") as zf:
            zf.extractall(extract_dir)
    except zipfile.BadZipFile:
        upload_path.unlink(missing_ok=True)
        shutil.rmtree(extract_dir, ignore_errors=True)
        return jsonify({"error": "Invalid ZIP file"}), 400

    docs = collect_documents(extract_dir)

    if not docs:
        shutil.rmtree(extract_dir, ignore_errors=True)
        upload_path.unlink(missing_ok=True)
        return jsonify({"error": "No readable documents found in ZIP"}), 400

    # Persist doc list for this session
    manifest_path = EXTRACTED_FOLDER / f"{session_id}_manifest.json"
    with open(manifest_path, "w") as f:
        json.dump(docs, f)

    file_list = [d["filename"] for d in docs]
    return jsonify({
        "session_id": session_id,
        "document_count": len(docs),
        "files": file_list,
    })


@app.route("/search", methods=["POST"])
def search():
    data = request.get_json()
    session_id = data.get("session_id", "")
    query = data.get("query", "").strip()

    if not session_id or not query:
        return jsonify({"error": "session_id and query are required"}), 400

    manifest_path = EXTRACTED_FOLDER / f"{session_id}_manifest.json"
    if not manifest_path.exists():
        return jsonify({"error": "Session not found. Please upload documents again."}), 404

    with open(manifest_path) as f:
        docs = json.load(f)

    context = build_document_context(docs)

    system_prompt = """You are a specialist in elevator and lift engineering specifications.
Your role is to analyse technical documents and extract relevant information about lift specifications.

When answering queries, structure your response as:
1. **Direct Answer** – concise summary of findings
2. **Specification Details** – specific values, standards, codes, or requirements found (use a table or list)
3. **Source References** – which document(s) contained the information
4. **Additional Notes** – related specifications or compliance considerations the user should be aware of

If the information is not found in the documents, clearly state that and suggest what to look for.
Always cite the specific document filename when referencing data."""

    user_message = f"""The following documents are lift/elevator specification documents that have been uploaded:

{context}

---

User query: {query}

Please search through these documents and provide a detailed, structured response specific to lift/elevator specifications."""

    try:
        message = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=2048,
            system=system_prompt,
            messages=[{"role": "user", "content": user_message}],
        )
        answer = message.content[0].text
        return jsonify({"answer": answer, "document_count": len(docs)})
    except anthropic.APIError as e:
        return jsonify({"error": f"AI service error: {e}"}), 500


@app.route("/session/<session_id>/files", methods=["GET"])
def list_files(session_id: str):
    manifest_path = EXTRACTED_FOLDER / f"{session_id}_manifest.json"
    if not manifest_path.exists():
        return jsonify({"error": "Session not found"}), 404
    with open(manifest_path) as f:
        docs = json.load(f)
    return jsonify({"files": [d["filename"] for d in docs]})


@app.route("/session/<session_id>", methods=["DELETE"])
def delete_session(session_id: str):
    extract_dir = EXTRACTED_FOLDER / session_id
    manifest_path = EXTRACTED_FOLDER / f"{session_id}_manifest.json"
    upload_path = UPLOAD_FOLDER / f"{session_id}.zip"

    shutil.rmtree(extract_dir, ignore_errors=True)
    manifest_path.unlink(missing_ok=True)
    upload_path.unlink(missing_ok=True)

    return jsonify({"message": "Session deleted"})


if __name__ == "__main__":
    UPLOAD_FOLDER.mkdir(exist_ok=True)
    EXTRACTED_FOLDER.mkdir(exist_ok=True)
    app.run(debug=True, port=5000)
